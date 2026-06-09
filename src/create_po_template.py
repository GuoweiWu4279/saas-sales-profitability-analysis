#!/usr/bin/env python3
"""
Creates a blank monthly PO tracking Excel with:
  - Sheet "MM-YYYY": data entry with auto-reviewer lookup formula
  - Sheet "Program Reviewers": your code→reviewer mapping
  - Sheet "Sheet2": spare

Run once to create the template, then fill in your PO data each month.

Usage:
    python src/create_po_template.py
    # Creates: PO_Tracking_06-2026.xlsx  (uses current month)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

# ---------------------------------------------------------------------------
# Paste your full program→reviewer list here (keep this updated)
# ---------------------------------------------------------------------------
REVIEWER_MAPPING = [
    # Program Code              Reviewer Name
    ("ACFlexMHSA01.1",          "Kanwar Kular"),
    ("ACFlexMHSA01.2",          "Kanwar Kular"),
    ("ACFlexBHSM01",            "Kanwardeep"),
    ("ACBFHome01",              "Kanwar Kular"),
    ("GH001",                   "Kanwardeep"),
    ("GH001.5",                 "Kanwardeep"),
    ("GH001.6",                 "Kanwardeep"),
    # Add more rows here as needed
]


def create_template():
    month_str = datetime.now().strftime("%m-%Y")   # e.g. "06-2026"
    filename = f"PO_Tracking_{month_str}.xlsx"

    wb = openpyxl.Workbook()

    # -----------------------------------------------------------------------
    # Sheet 1: monthly PO data
    # -----------------------------------------------------------------------
    ws_po = wb.active
    ws_po.title = month_str

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["PO Number", "Program Code", "PO, Program", "Reviewer", "Status", "Notes"]
    col_widths = [12, 22, 30, 20, 12, 30]

    for col, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws_po.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws_po.column_dimensions[get_column_letter(col)].width = w

    # Pre-fill formulas for rows 2–200
    for row in range(2, 201):
        a = f"A{row}"
        b = f"B{row}"
        # Column C: combined label
        ws_po[f"C{row}"] = f'=IF(AND({a}<>"",{b}<>""),{a}&", "&{b},"")'
        # Column D: auto-lookup reviewer from Program Reviewers sheet
        ws_po[f"D{row}"] = (
            f'=IF({b}="","",IFERROR(VLOOKUP({b},\'Program Reviewers\'!$A:$B,2,0),"no result"))'
        )

    ws_po.freeze_panes = "A2"

    # -----------------------------------------------------------------------
    # Sheet 2: Program Reviewers mapping
    # -----------------------------------------------------------------------
    ws_map = wb.create_sheet("Program Reviewers")
    map_headers = ["Program Code", "Reviewer Name"]
    for col, h in enumerate(map_headers, start=1):
        cell = ws_map.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws_map.column_dimensions["A"].width = 25
    ws_map.column_dimensions["B"].width = 22

    for row, (code, reviewer) in enumerate(REVIEWER_MAPPING, start=2):
        ws_map.cell(row=row, column=1, value=code)
        ws_map.cell(row=row, column=2, value=reviewer)

    ws_map.freeze_panes = "A2"

    # -----------------------------------------------------------------------
    # Sheet 3: spare
    # -----------------------------------------------------------------------
    wb.create_sheet("Sheet2")

    # -----------------------------------------------------------------------
    # Save
    # -----------------------------------------------------------------------
    output_path = os.path.join("output", filename)
    os.makedirs("output", exist_ok=True)
    wb.save(output_path)
    print(f"Template created: {output_path}")
    print("  → Open it, fill Column A (PO#) and Column B (Program Code).")
    print("    Column D (Reviewer) fills automatically via VLOOKUP.")
    print(f"  → Then run: python src/po_email_generator.py {output_path}")


if __name__ == "__main__":
    create_template()
